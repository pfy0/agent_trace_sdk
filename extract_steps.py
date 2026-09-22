"""将 messages 数据提取为每一步一行的 Excel 表。

输入支持多层嵌套结构（如 {"conversations": {"messages": [...]}}）、
{"messages": [...]} 以及裸数组。
输出列: step_no, message_id, role, message_type, content,
        timestamp, tool_calls, tool_call_id, latency_ms

一条 tool_call 消息若包含多个工具调用，会按调用拆成多行：
每行保留相同的 message_id，tool_calls 只存对应的那一个调用对象，
step_no 依次递增。tool_call_id 同步填入该调用的 id，便于与
tool_response 一一对齐。
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

MAX_DEPTH = 8

# message_type -> (类型缩写, 对应 role)
TYPE_MAP = {
    "system_prompt": ("sys", "system"),
    "user_query": ("usr", "user"),
    "reasoning": ("rsn", "assistant"),
    "tool_call": ("cal", "assistant"),
    "tool_response": ("res", "tool"),
    "final_answer": ("ans", "assistant"),
}

COLUMNS = [
    "step_no",
    "message_id",
    "role",
    "message_type",
    "content",
    "timestamp",
    "tool_calls",
    "tool_call_id",
    "latency_ms",
]


def find_messages(node, depth=0):
    """在任意嵌套结构中定位 messages 列表，找不到返回 None。"""
    if depth > MAX_DEPTH:
        return None
    if isinstance(node, list):
        if node and all(isinstance(item, dict) and "message_type" in item for item in node):
            return node
        return None
    if isinstance(node, dict):
        inner = node.get("messages")
        if isinstance(inner, list):
            return inner
        for value in node.values():
            found = find_messages(value, depth + 1)
            if found is not None:
                return found
    return None


def load_messages(path):
    """读取消息列表；结构不对时直接报错，避免静默产出空表。"""
    with path.open(encoding="utf-8") as fh:
        data = json.load(fh)

    messages = find_messages(data)
    if messages is None:
        raise ValueError("未在 %s 中找到 messages 列表，请检查 JSON 结构" % path)
    if not messages:
        raise ValueError("%s 中的 messages 列表为空，没有可导出的数据" % path)
    return messages


def to_cell(value):
    """将任意值转为可写入单元格的值，复杂结构序列化为 JSON 字符串。"""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2)


def build_rows(messages):
    """按消息顺序生成表格行；tool_call 多调用时拆行，step_no 从 1 递增。"""
    rows = []

    def emit(message, tool_calls, tool_call_id):
        message_type = message.get("message_type")
        rows.append(
            {
                "step_no": len(rows) + 1,
                "message_id": message.get("message_id"),
                "role": message.get("role") or TYPE_MAP[message_type][1],
                "message_type": message_type,
                "content": to_cell(message.get("content")),
                "timestamp": message.get("timestamp"),
                "tool_calls": to_cell(tool_calls),
                "tool_call_id": tool_call_id,
                "latency_ms": message.get("latency_ms"),
            }
        )

    for index, message in enumerate(messages, start=1):
        message_type = message.get("message_type")
        if message_type not in TYPE_MAP:
            raise ValueError("第 %d 条消息的 message_type 未知: %r" % (index, message_type))

        if message_type == "tool_call":
            tool_calls = message.get("tool_calls") or []
            if not tool_calls:
                raise ValueError("第 %d 条消息为 tool_call 但缺少 tool_calls" % index)
            for tool_call in tool_calls:
                emit(message, tool_call, tool_call.get("id"))
            continue

        if message_type == "tool_response" and not message.get("tool_call_id"):
            raise ValueError("第 %d 条消息为 tool_response 但缺少 tool_call_id" % index)

        emit(message, message.get("tool_calls"), message.get("tool_call_id"))

    return rows


def write_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise SystemExit("缺少依赖 openpyxl，请先执行: pip install openpyxl") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "messages"

    sheet.append(COLUMNS)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for column_index in range(1, len(COLUMNS) + 1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row[column] for column in COLUMNS])

    widths = {
        "step_no": 8,
        "message_id": 14,
        "role": 11,
        "message_type": 14,
        "content": 70,
        "timestamp": 25,
        "tool_calls": 55,
        "tool_call_id": 28,
        "latency_ms": 11,
    }
    for column_index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = widths[column]

    wrap = Alignment(wrap_text=True, vertical="top")
    for row_index in range(2, sheet.max_row + 1):
        for column_index in range(1, len(COLUMNS) + 1):
            sheet.cell(row=row_index, column=column_index).alignment = wrap

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description="提取 messages 为 Excel 步骤表")
    parser.add_argument("input", nargs="?", default="data/sample_messages.json", help="输入 JSON 文件")
    parser.add_argument("-o", "--output", default="output/messages_steps.xlsx", help="输出 xlsx 路径")
    args = parser.parse_args()

    rows = build_rows(load_messages(Path(args.input)))
    write_xlsx(rows, Path(args.output))
    print("已写入 %d 行 -> %s" % (len(rows), args.output))


if __name__ == "__main__":
    main()
