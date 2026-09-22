"""XLSX 导出器：把 trace 的 conversations 展平成一行一个 step 的 Excel 表。

一条 tool_call 消息 = 一行，多个调用放在同一格里（JSON 数组）。
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Union

from ..schemas import MessageType, Trace
from .base import Exporter

COLUMNS = [
    "step_no",
    "message_id",
    "role",
    "message_type",
    "content",
    "timestamp",
    "tool_calls",
    "tool_call_id",
    "latency_ms",
]

COLUMN_WIDTHS = {
    "step_no": 8,
    "message_id": 14,
    "role": 11,
    "message_type": 14,
    "content": 70,
    "timestamp": 25,
    "tool_calls": 55,
    "tool_call_id": 28,
    "latency_ms": 11,
}


def to_cell(value: Any) -> Any:
    """复杂结构序列化为 JSON 字符串，其余原样返回。"""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2)


def trace_to_rows(trace: Trace) -> List[Dict[str, Any]]:
    """把一条 trace 展平成表格行。step_no 从 1 重新排，保持连续。"""
    rows: List[Dict[str, Any]] = []

    def emit(step: Any, tool_calls: Any, tool_call_id: Any) -> None:
        rows.append(
            {
                "step_no": len(rows) + 1,
                "message_id": step.message_id,
                "role": step.role.value,
                "message_type": step.message_type.value,
                "content": to_cell(step.content),
                "timestamp": step.timestamp,
                "tool_calls": to_cell(tool_calls),
                "tool_call_id": tool_call_id,
                "latency_ms": step.latency_ms,
            }
        )

    for step in trace.conversations:
        if step.message_type is MessageType.TOOL_CALL:
            calls = [call.model_dump(mode="json", exclude_none=True) for call in step.tool_calls or []]
            emit(step, calls, step.tool_call_id)
            continue
        emit(step, None, step.tool_call_id)

    return rows


def write_xlsx(rows: List[Dict[str, Any]], path: Union[str, Path]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ImportError("缺少依赖 openpyxl，请先执行: pip install openpyxl") from exc

    path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "messages"

    sheet.append(COLUMNS)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for column_index in range(1, len(COLUMNS) + 1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row[column] for column in COLUMNS])

    for column_index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = COLUMN_WIDTHS[column]

    wrap = Alignment(wrap_text=True, vertical="top")
    for row_index in range(2, sheet.max_row + 1):
        for column_index in range(1, len(COLUMNS) + 1):
            sheet.cell(row=row_index, column=column_index).alignment = wrap

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


class XlsxExporter(Exporter):
    """每次 export 把累积的行重写进同一个 xlsx。"""

    def __init__(
        self,
        path: Union[str, Path],
        flush_every: int = 1,
    ) -> None:
        self.path = Path(path)
        self.flush_every = max(1, int(flush_every))
        self._rows: List[Dict[str, Any]] = []
        self._pending = 0

    def export(self, trace: Trace) -> None:
        self._rows.extend(trace_to_rows(trace))
        self._pending += 1
        if self._pending >= self.flush_every:
            self.flush()

    def flush(self) -> None:
        self._pending = 0
        write_xlsx(self._rows, self.path)

    @property
    def rows(self) -> List[Dict[str, Any]]:
        return list(self._rows)
